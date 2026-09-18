# -*- coding: utf-8 -*-
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ST = json.load(open("state.json", encoding="utf-8"))
S=[{"tr":sec["title"]["tr"],"g":[{"tr":g["title"]["tr"],"i":[{"n":it["name"]["tr"],"ne":it["name"]["en"],"p":it["p"],"ing":it["ing"]["tr"],"inge":it["ing"]["en"],"a":it["a"],"kcal":it.get("kcal"),"nut":it.get("nut"),"sos":it.get("sos"),"alc":it.get("alc",0),"pork":it.get("pork",0),"sec_id":sec["id"]} for it in g["i"]]} for g in sec["g"]]} for sec in ST["sections"]]
ALG = {"G":"Gluten","M":"Süt","Y":"Yumurta","B":"Balık","K":"Kabuklu deniz ürünleri",
       "S":"Soya","SS":"Susam","N":"Sert kabuklu yemişler","F":"Yer fıstığı",
       "H":"Hardal","C":"Kereviz","SO":"Sülfit","MO":"Yumuşakçalar","L":"Acı bakla"}
CODES = list(ALG.keys())

INK="FF21201A"; MOSS="FF1A1814"; BRASS="FF8A6A1F"
wb = Workbook(); ws = wb.active; ws.title = "Menü"

hdr = ["Bölüm","Grup","Ürün (TR)","Ürün (EN)","Fiyat ₺","İçindekiler (TR)","İçindekiler (EN)"] + CODES + ["Enerji kcal (porsiyon)","Protein g","Yağ g","Karbonhidrat g","Alkol (X)","Domuz türevi (X)","Sos kartı","Notunuz"]
ws.append(hdr)

thin = Side(style="thin", color="FFD5D2C4")
for c in range(1, len(hdr)+1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFFFF")
    cell.fill = PatternFill("solid", fgColor=MOSS)
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
ws.row_dimensions[1].height = 40

r = 2
for s in S:
    for g in s["g"]:
        for it in g["i"]:
            row = [s["tr"], g["tr"], it["n"], it["ne"], it["p"],
                   it.get("ing",""), it.get("inge","")]
            a = set(it.get("a", []))
            row += ["X" if c in a else "" for c in CODES]
            nu = it.get("nut") or {}
            row += [("" if it.get("kcal") is None else it["kcal"]), nu.get("p",""), nu.get("f",""), nu.get("c",""), ("X" if it.get("alc") else ""), ("X" if it.get("pork") else ""), ", ".join(it.get("sos") or []), ""]
            ws.append(row)
            for c in range(1, len(hdr)+1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(name="Arial", size=10)
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=(c in (6,7)))
            ws.cell(row=r, column=5).number_format = '#,##0'
            ws.cell(row=r, column=5).font = Font(name="Arial", size=10, bold=True, color=BRASS)
            ws.cell(row=r, column=5).alignment = Alignment(vertical="top", horizontal="right")
            for c in range(8, 8+len(CODES)):
                cell = ws.cell(row=r, column=c)
                cell.alignment = Alignment(vertical="center", horizontal="center")
                cell.font = Font(name="Arial", size=10, bold=True, color=BRASS)
                cell.fill = PatternFill("solid", fgColor="FFFFF9E8")
            for c in range(8+len(CODES), len(hdr)+1):
                cell = ws.cell(row=r, column=c)
                cell.fill = PatternFill("solid", fgColor="FFFFFFCC")
                cell.alignment = Alignment(vertical="center", horizontal="center" if c < len(hdr) else "left")
            for c in range(8+len(CODES), 12+len(CODES)):
                ws.cell(row=r, column=c).font = Font(name="Arial", size=10, bold=(c==8+len(CODES)), color=BRASS)
            r += 1
last = r-1

widths = [15,24,32,32,9,52,52] + [5]*len(CODES) + [12,9,8,13,9,11,18,26]
for i,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{last}"

dv = DataValidation(type="list", formula1='"X,"', allow_blank=True, showDropDown=False)
ws.add_data_validation(dv)
dv.add(f"H2:{get_column_letter(7+len(CODES))}{last}")
dv.add(f"{get_column_letter(12+len(CODES))}2:{get_column_letter(13+len(CODES))}{last}")

# --- Nasıl kullanılır ---
ws2 = wb.create_sheet("Nasıl kullanılır")
lines = [
 ("KULE CAFE — Alerjen, içerik ve enerji kontrol listesi", True),
 ("", False),
 ("İçindekiler: mutfağın 16.09.2026 tarihli düzeltmeleri işlenmiştir.", False),
 ("Alerjen işaretleri ve BESİN DEĞERLERİ (kcal, protein, yağ, karbonhidrat) TASLAKTIR: standart tarif ve porsiyon", False),
 ("büyüklüklerinden hesaplanmış yaklaşık değerlerdir. Mevzuata uygun beyan için mutfak", False),
 ("gramajlarıyla bir gıda mühendisi / diyetisyen tarafından doğrulanmalıdır.", False),
 ("Alkol ve Domuz türevi sütunları: ürün içeriyorsa X yazın (şu an hepsi boş = içermez).", False),
 ("Sos kartı sütunu: o üründe hangi hazır sosun ayrı kartı gösteriliyor (sweetchili / cafedeparis / demiglace).", False),
 ("", False),
 ("Nasıl düzeltirsiniz:", True),
 ("1. \"Menü\" sekmesinde her satır bir üründür.", False),
 ("2. Alerjen sütunlarında (G, M, Y, ...) ürün o alerjeni içeriyorsa hücreye X yazın,", False),
 ("   içermiyorsa hücreyi boş bırakın. Sarı hücreler doldurulacak alanlardır.", False),
 ("3. İçindekiler sütunlarını serbestçe düzeltebilir, boş olanları doldurabilirsiniz.", False),
 ("4. Enerji / protein / yağ / karbonhidrat sütunlarına porsiyon başına değerleri yazın; Alkol / Domuz türevi sütunlarına içeriyorsa X.", False),
 ("5. Fiyat değişikliklerini de bu dosyada yapabilirsiniz.", False),
 ("6. Dosyayı bana geri gönderin; menüyü güncellerim. QR kod değişmez.", False),
 ("", False),
 ("Örnek satır:", True),
 ("Bölüm: Tatlılar | Grup: Tatlılar | Ürün: Künefe | Fiyat: 420", False),
 ("İçindekiler: Kadayıf, dil peyniri, tereyağı, şerbet, Antep fıstığı", False),
 ("G: X   M: X   N: X   (gluten, süt ve sert kabuklu yemiş içerir)", False),
 ("", False),
 ("Alerjen kodları:", True),
]
for code, name in ALG.items():
    lines.append((f"   {code}  =  {name}", False))
for i,(txt,bold) in enumerate(lines, start=1):
    c = ws2.cell(row=i, column=1, value=txt)
    c.font = Font(name="Arial", size=11, bold=bold, color=MOSS if bold else INK)
ws2.column_dimensions["A"].width = 95

wb.save("KULE-alerjen-kontrol.xlsx")
print("satır:", last-1, "| sütun:", len(hdr))
